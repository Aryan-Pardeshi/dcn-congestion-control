"""Build docs/Experiment_Congestion_Control.docx from content.json + observations.json + PNG."""
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

HERE = Path(__file__).resolve().parent
PROJ = HERE.parent
DOCS = PROJ.parent / "docs"
CONTENT = DOCS / "content.json"
OBS = DOCS / "observations.json"
PNG = DOCS / "network-diagram.png"
OUT = DOCS / "Experiment_Congestion_Control.docx"

HEAD_COLOR = RGBColor(0x12, 0x30, 0x4D)
HDR_SHADE = "DCE6EF"
CODE_SHADE = "EAEFF4"


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "8AA0B8")
        borders.append(el)
    tcPr.append(borders)


def bordered(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell)
            for p in cell.paragraphs:
                p.style.font.size = Pt(10)
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.name = "Calibri"


def header_row(row):
    for cell in row.cells:
        shade(cell, HDR_SHADE)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True


def para(doc, text, bold=False, size=11, color=None, align=None, font="Calibri"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    if bold:
        r.bold = True
    if color is not None:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return p


def h1(doc, text):
    h = doc.add_heading(text, level=1)
    for r in h.runs:
        r.font.color.rgb = HEAD_COLOR
        r.bold = True
    return h


def h2(doc, text):
    h = doc.add_heading(text, level=2)
    for r in h.runs:
        r.font.color.rgb = HEAD_COLOR
        r.bold = True
    return h


def bullets(doc, items, numbered=False):
    for it in items:
        doc.add_paragraph(it, style="List Number" if numbered else "List Bullet")


def code_block(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, CODE_SHADE)
    set_cell_borders(cell)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    return tbl


def add_page_number(footer_para):
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_para.add_run("Page ")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    f2 = OxmlElement("w:instrText")
    f2.set(qn("xml:space"), "preserve")
    f2.text = "PAGE"
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    for el in (f1, f2, f3):
        rr = OxmlElement("w:r")
        rr.append(el)
        footer_para._p.append(rr)


def main():
    content = json.loads(CONTENT.read_text(encoding="utf-8"))
    obs = json.loads(OBS.read_text(encoding="utf-8"))
    exp = content["experiment"]
    ps = content["problemStatement"]
    res = content["resources"]
    cfg = content["DEFAULT_CONFIG"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # header / footer
    hp = sec.header.paragraphs[0]
    hr = hp.add_run("DCN Lab \u2014 Congestion Control")
    hr.font.size = Pt(9)
    hr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    add_page_number(sec.footer.paragraphs[0])

    # title block
    para(doc, "Experiment: Implementation of Congestion Control Algorithms",
         bold=True, size=16, color=HEAD_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"{exp['module']}  |  {exp['unit']}", size=11,
         color=RGBColor(0x55, 0x55, 0x55), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Name: ______    Roll no.: ______    Date: ______", size=11,
         align=WD_ALIGN_PARAGRAPH.CENTER)

    # 1 Aim
    h1(doc, "1  Aim")
    para(doc, exp["aim"])

    # 2 Problem statement
    h1(doc, "2  Problem statement")
    for s in ps["scenario"]:
        para(doc, s)
    h2(doc, "Network requirements")
    bullets(doc, ps["requirements"])
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "Metric", "Target"
    header_row(tbl.rows[0])
    for o in ps["objectives"]:
        row = tbl.add_row().cells
        row[0].text, row[1].text = o["metric"], o["target"]
    bordered(tbl)

    # 3 Resources
    h1(doc, "3  Resources needed")
    h2(doc, "Hardware")
    bullets(doc, res["hardware"])
    h2(doc, "Software")
    bullets(doc, res["software"])
    h2(doc, "References")
    bullets(doc, res["references"], numbered=True)

    # 4 Theory
    h1(doc, "4  Theory")
    for block in content["theoryIntro"]:
        h2(doc, block["heading"])
        for b in block["body"]:
            para(doc, b)
    h2(doc, "Mechanisms used to control congestion")
    para(doc, "Leaky bucket and token bucket are traffic shaping mechanisms (open loop, applied at the "
              "network entry). The choke packet is a feedback congestion control mechanism (closed loop, "
              "router signals the source). RED is an active queue management mechanism (inside the router queue).")
    for m in content["mechanisms"]:
        h2(doc, f"{m['name']} \u2014 {m['category']}")
        para(doc, f"Category: {m['category']}  |  Loop type: {m['loop']}  |  Acts at: {m['where']}")
        para(doc, m["summary"])
        para(doc, "How it works:", bold=True)
        bullets(doc, m["howItWorks"], numbered=True)
        p = para(doc, "Formula: ", bold=True)
        fr = p.add_run(m["formula"])
        fr.font.name = "Consolas"
        fr.bold = False
        para(doc, "Advantages:", bold=True)
        bullets(doc, m["pros"])
        para(doc, "Limitations:", bold=True)
        bullets(doc, m["cons"])

    comp = content["comparison"]
    para(doc, "Comparison of mechanisms:", bold=True)
    ct = doc.add_table(rows=1, cols=len(comp["columns"]))
    ct.style = "Table Grid"
    for i, c in enumerate(comp["columns"]):
        ct.rows[0].cells[i].text = c
    header_row(ct.rows[0])
    for row in comp["rows"]:
        cells = ct.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    bordered(ct)

    # 5 Network design
    h1(doc, "5  Network design")
    doc.add_picture(str(PNG), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Figure 1: Campus network with bottleneck WAN link at gateway router R1",
         size=10, color=RGBColor(0x55, 0x55, 0x55), align=WD_ALIGN_PARAGRAPH.CENTER)
    heads = ["Device", "Interface", "IP address", "Subnet", "Link"]
    at = doc.add_table(rows=1, cols=len(heads))
    at.style = "Table Grid"
    for i, hh in enumerate(heads):
        at.rows[0].cells[i].text = hh
    header_row(at.rows[0])
    for a in content["addressing"]:
        cells = at.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text, cells[4].text = (
            a["device"], a["interface"], a["ip"], a["subnet"], a["link"])
    bordered(at)

    # 6 Simulation set-up
    h1(doc, "6  Simulation set-up")
    para(doc, "Simulations were run with the custom Vite + React web simulator in this repository "
              "(npm run dev). Each mechanism was run on identical traffic (same seed) for 1000 ticks "
              "using compareAll(DEFAULT_CONFIG, 1000) so the results are directly comparable.")
    st = doc.add_table(rows=1, cols=2)
    st.style = "Table Grid"
    st.rows[0].cells[0].text, st.rows[0].cells[1].text = "Parameter", "Value"
    header_row(st.rows[0])
    srows = [
        ("Link capacity (R1 service rate, packets/tick)", str(cfg["linkCapacity"])),
        ("R1 buffer size (packets)", str(cfg["bufferSize"])),
        ("Random seed", str(cfg["seed"])),
    ]
    for s in cfg["sources"]:
        srows.append((f"Source {s['id']} ({s['label']}): rate, pattern", f"{s['rate']}, {s['pattern']}"))
    for mech, params in cfg["params"].items():
        for k, v in params.items():
            srows.append((f"{mech}.{k}", str(v)))
    for k, v in srows:
        cells = st.add_row().cells
        cells[0].text, cells[1].text = k, v
    bordered(st)

    # 7 Procedure
    h1(doc, "7  Procedure")
    bullets(doc, content["procedure"], numbered=True)

    # 8 Observations
    h1(doc, "8  Observations")
    oh = ["Mechanism", "Category", "Generated", "Delivered", "Dropped",
          "Drop rate %", "Avg queue", "Avg delay", "Throughput", "Utilisation %"]
    ot = doc.add_table(rows=1, cols=len(oh))
    ot.style = "Table Grid"
    for i, hh in enumerate(oh):
        ot.rows[0].cells[i].text = hh
    header_row(ot.rows[0])
    for r in obs:
        cells = ot.add_row().cells
        vals = [r["name"], r.get("category", ""),
                str(r["generated"]), str(r["delivered"]), str(r["dropped"]),
                f"{r['dropRate'] * 100:.2f}", f"{r['avgQueue']:.2f}",
                f"{r['avgDelay']:.2f}", f"{r['throughput']:.2f}",
                f"{r['utilization'] * 100:.2f}"]
        for i, v in enumerate(vals):
            cells[i].text = v
    bordered(ot)

    by_drop = min(obs, key=lambda r: r["dropRate"])
    by_queue = min(obs, key=lambda r: r["avgQueue"])
    by_delay = min(obs, key=lambda r: r["avgDelay"])
    base = next((r for r in obs if r["id"] == "none"), obs[0])
    others = [r for r in obs if r["id"] != base["id"]]
    best_ctrl = min(others, key=lambda r: r["dropRate"]) if others else base
    imp = (base["dropRate"] - best_ctrl["dropRate"]) * 100 if others else 0.0
    para(doc, f"The lowest packet drop rate was achieved by {by_drop['name']} "
              f"({by_drop['dropRate'] * 100:.2f}%), compared with {base['dropRate'] * 100:.2f}% "
              f"for the no-control baseline.")
    para(doc, f"The smallest average queue ({by_queue['avgQueue']:.2f} packets) was recorded under "
              f"{by_queue['name']}, and the lowest average queuing delay ({by_delay['avgDelay']:.2f} ticks) "
              f"under {by_delay['name']}.")
    para(doc, f"Without any control the bottleneck queue stays nearly full (average queue "
              f"{base['avgQueue']:.2f} packets) and {base['dropped']} of {base['generated']} packets are lost, "
              f"whereas the best controlled run ({best_ctrl['name']}) cuts the drop rate by "
              f"{imp:.2f} percentage points while keeping utilisation at "
              f"{best_ctrl['utilization'] * 100:.2f}%.")
    lower_delay = [r["name"] for r in others if r["avgDelay"] < base["avgDelay"]]
    lower_queue = [r["name"] for r in others if r["avgQueue"] < base["avgQueue"]]
    para(doc, "Compared with the baseline, a lower average queue was seen with "
              f"{', '.join(lower_queue) or 'none of the mechanisms'}, and a lower average delay with "
              f"{', '.join(lower_delay) or 'none of the mechanisms'}.")
    para(doc, "Because the offered load stays above the bottleneck capacity, the open-loop mechanisms "
              "(traffic shaping) and RED cannot remove the excess traffic: they only decide where and when "
              "it is dropped, so their drop rate stays close to the baseline. Their benefit is a smoother "
              "input (leaky bucket), controlled bursts (token bucket) or a shorter queue and lower delay (RED). "
              "Only the closed-loop choke packet mechanism lowers the offered load itself, which is why its "
              "drop rate falls the most. With TCP-like senders that slow down after loss, RED's early drops "
              "would also reduce the offered load.")

    # 9 Result
    h1(doc, "9  Result")
    para(doc, "The campus network was designed with addressing for all LANs and simulated with all "
              "congestion control mechanisms on identical traffic. "
              f"The best drop-rate performer was {best_ctrl['name']} "
              f"({best_ctrl['dropRate'] * 100:.2f}% drops), while {by_delay['name']} gave the lowest "
              f"average delay ({by_delay['avgDelay']:.2f} ticks), against "
              f"{base['dropRate'] * 100:.2f}% drops and {base['avgDelay']:.2f} ticks delay for the "
              "no-control baseline.")

    # 10 Conclusion
    h1(doc, "10  Conclusion")
    for c in content["conclusion"]:
        para(doc, c)

    # 11 Viva
    h1(doc, "11  Viva questions")
    bullets(doc, content["vivaQuestions"], numbered=True)

    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
